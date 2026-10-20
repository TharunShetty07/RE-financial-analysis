
from __future__ import annotations

from datetime import datetime
from pathlib import Path
from typing import Any

import openpyxl


MONTH_FORMATS = ("%b %Y", "%d %b %Y")

INCOME_TOTAL_ROWS = {
    "Total Income",
    "Total Cost of Goods Sold",
    "Gross Profit",
    "Total S & M Payroll",
    "Total Marketing",
    "Total Operating Expenses",
    "Operating Income",
    "Net Income",
}

BALANCE_TOTAL_ROWS = {
    "Total Cash and Cash Equivalents",
    "Total Current Assets",
    "Total Fixed Assets",
    "Total Long Term Assets",
    "Total Assets",
    "Total Current Liabilities",
    "Total Long Term Liabilities",
    "Total Liabilities",
    "Total Equity",
    "Total Liabilities and Equity",
}


def parse_month(value: Any) -> datetime:
    """Parse a supported month/date header into a sortable date."""
    if isinstance(value, datetime):
        return datetime(value.year, value.month, 1)

    if isinstance(value, str):
        for fmt in MONTH_FORMATS:
            try:
                parsed = datetime.strptime(value.strip(), fmt)
                return datetime(parsed.year, parsed.month, 1)
            except ValueError:
                continue

    raise ValueError(f"Unsupported month header: {value!r}")


def find_account_label(row: tuple[Any, ...]) -> str | None:
    """Return the first non-empty text cell in a worksheet row."""
    for value in row:
        if isinstance(value, str) and value.strip():
            return value.strip()
    return None


def load_workbook_data(
    file_path: str | Path,
) -> dict[str, Any]:
    """
    Load monthly financial statements without modifying the source file.

    Returns chronological months, account values, reported totals,
    and validation warnings.
    """
    path = Path(file_path)

    if not path.is_file():
        raise FileNotFoundError(f"Workbook not found: {path}")

    workbook = openpyxl.load_workbook(
        path,
        data_only=True,
        read_only=False,
    )

    try:
        required_sheets = {"Income Statement", "Balance Sheet"}
        missing_sheets = required_sheets - set(workbook.sheetnames)

        if missing_sheets:
            raise ValueError(
                f"Missing required worksheets: {sorted(missing_sheets)}"
            )

        result: dict[str, Any] = {
            "source_file": path.name,
            "months": [],
            "income_statement": {},
            "balance_sheet": {},
            "reported_totals": {},
            "warnings": [],
        }

        sheet_specs = {
            "Income Statement": {
                "header_row": 5,
                "first_month_col": 3,
                "output_key": "income_statement",
                "total_labels": INCOME_TOTAL_ROWS,
            },
            "Balance Sheet": {
                "header_row": 5,
                "first_month_col": 4,
                "output_key": "balance_sheet",
                "total_labels": BALANCE_TOTAL_ROWS,
            },
        }

        canonical_months: list[datetime] | None = None

        for sheet_name, spec in sheet_specs.items():
            sheet = workbook[sheet_name]
            header_row = spec["header_row"]
            first_month_col = spec["first_month_col"]

            headers = [
                sheet.cell(header_row, col).value
                for col in range(
                    first_month_col,
                    sheet.max_column + 1,
                )
            ]

            month_columns = []
            for col, header in enumerate(
                headers,
                start=first_month_col,
            ):
                if header is None:
                    continue
                try:
                    month_columns.append((parse_month(header), col))
                except ValueError:
                    # Ignore non-date columns, but don't guess their meaning.
                    continue

            if not month_columns:
                raise ValueError(
                    f"No monthly columns found in {sheet_name}"
                )

            month_columns.sort(key=lambda item: item[0])
            months = [month for month, _ in month_columns]

            if len(set(months)) != len(months):
                raise ValueError(
                    f"Duplicate monthly headers in {sheet_name}"
                )

            if canonical_months is None:
                canonical_months = months
            elif months != canonical_months:
                raise ValueError(
                    "The two statements do not cover the same months."
                )

            accounts: dict[str, dict[str, float | None]] = {}
            totals: dict[str, dict[str, float | None]] = {}
            total_labels = spec["total_labels"]

            for row_num in range(header_row + 1, sheet.max_row + 1):
                row = tuple(
                    sheet.cell(row_num, col).value
                    for col in range(1, sheet.max_column + 1)
                )
                label = find_account_label(row)

                if not label:
                    continue

                # Section headings are not account data.
                if label in {
                    "Income",
                    "Cost of Goods Sold",
                    "Operating Expenses",
                    "Assets",
                    "Current Assets",
                    "Fixed Assets",
                    "Long Term Assets",
                    "Liabilities and Equity",
                    "Liabilities",
                    "Current Liabilities",
                    "Long Term Liabilities",
                    "Equity",
                    "Cash and Cash Equivalents",
                    "S & M Payroll",
                    "Marketing",
                }:
                    continue

                values: dict[str, float | None] = {}

                for month, col in month_columns:
                    value = sheet.cell(row_num, col).value
                    month_key = month.strftime("%Y-%m")

                    if value is None:
                        values[month_key] = None
                    elif isinstance(value, bool):
                        raise ValueError(
                            f"Unexpected boolean at {sheet_name}!{col}{row_num}"
                        )
                    elif isinstance(value, (int, float)):
                        values[month_key] = float(value)
                    else:
                        result["warnings"].append(
                            f"{sheet_name}, row {row_num}, "
                            f"{label}, {month_key}: "
                            f"non-numeric value {value!r}"
                        )
                        values[month_key] = None

                destination = (
                    totals if label in total_labels else accounts
                )

                if label in destination:
                    result["warnings"].append(
                        f"{sheet_name}: duplicate account label "
                        f"{label!r}; row {row_num}"
                    )
                    label = f"{label} [row {row_num}]"

                destination[label] = values

                if label in total_labels:
                    result["reported_totals"].setdefault(
                        spec["output_key"], {}
                    )[label] = values

                if label not in total_labels:
                    missing_months = [
                        month for month, value in values.items()
                        if value is None
                    ]
                    if missing_months:
                        result["warnings"].append(
                            f"{sheet_name}, {label}: missing values "
                            f"for {', '.join(missing_months)}"
                        )

                    for month, value in values.items():
                        if value is not None and value < 0:
                            result["warnings"].append(
                                f"{sheet_name}, {label}, {month}: "
                                f"negative value {value:.2f}; "
                                "preserved for review"
                            )

            result[spec["output_key"]] = accounts
            result["reported_totals"][spec["output_key"]] = totals

        result["months"] = [
            month.strftime("%Y-%m")
            for month in (canonical_months or [])
        ]

        _validate_balance_sheet(result)
        return result

    finally:
        workbook.close()


def _validate_balance_sheet(data: dict[str, Any]) -> None:
    """Check reported assets against reported liabilities and equity."""
    totals = data["reported_totals"]["balance_sheet"]
    assets = totals.get("Total Assets", {})
    liabilities_equity = totals.get("Total Liabilities and Equity", {})

    for month in data["months"]:
        a = assets.get(month)
        le = liabilities_equity.get(month)

        if a is None or le is None:
            data["warnings"].append(
                f"{month}: cannot validate balance sheet; "
                "a required reported total is missing"
            )
            continue

        difference = a - le

        if abs(difference) > 0.01:
            data["warnings"].append(
                f"{month}: balance sheet does not balance; "
                f"assets minus liabilities and equity = {difference:.2f}"
            )
